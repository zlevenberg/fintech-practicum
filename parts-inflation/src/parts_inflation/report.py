"""Excel results workbook generation with professional formatting."""

from __future__ import annotations

import logging
import platform
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="D6DCE4")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, max_width: int = 36) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:50]:
            if cell.value is not None:
                length = max(length, min(max_width, len(str(cell.value))))
        ws.column_dimensions[letter].width = max(12, length + 2)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1) -> None:
    if df is None or df.empty:
        ws.cell(start_row, 1, "No data")
        return
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start_row):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.border = THIN
            if r_idx == start_row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
    ws.freeze_panes = f"A{start_row + 1}"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)


METHODOLOGY_TEXT = """
METHODOLOGY OVERVIEW

1. Scope
The model estimates client-experienced parts inflation from repeated purchases of the same PartKey.
Default scope is physical_inputs (inventory, production/operating supplies, production aids, small tools, etc.).
Labor, outside services, freight/shipping, and service-only repairs are excluded by default.
Needs Review categories are excluded from the default model but quantified in Scope Sensitivity.

2. Benchmarks (computed before the hierarchical model)
- Last-price: zero-inflation naive forecast using the latest observed price.
- Matched-part log-change index: median, equal-weight geometric mean, winsorized geometric mean,
  and capped-spend-weighted geometric mean of log(p_t / p_{t-1}) for parts observed in adjacent periods.
- Törnqvist index: expenditure-share-weighted average of matched log price changes with shares
  averaged across adjacent periods over the matched basket.
- Fisher index: geometric mean of Laspeyres and Paasche when quantities are valid; otherwise marked unavailable.

3. Repeat-purchase hierarchical model
Adjacent same-part purchases only (not all pairs). For interval j of part i:
  y_j = log(p_b / p_a)
  x_q,j = log(q_b / q_a)
  D_j,m = fraction of calendar month m in (d_a, d_b]
  y_j = sum_m D_j,m (delta0_m + u_c,m) + (gamma0 + kappa_c) x_q,j + e_j
Estimated with Huber IRLS and smoothness/ridge penalties on monthly rates and category deviations.
Quantity elasticity is shrunk toward zero and not forced to be negative.

4. Part hierarchy / fallback
1) Qualifying part-specific category + shrunk residual
2) Approved category forecast
3) Overall physical-input forecast
Shrinkage: lambda = n_eff/(n_eff+k) * min(span/730,1) * quality

5. Forecasting
Candidate future monthly log rates: trailing-12m mean, EWMA, damped Holt, mean reversion.
Selected via rolling-origin backtests on historical monthly rates.

6. Composite multiplier
Fixed-basket weights from trailing-12-month PO Value (or Planned Basket if supplied):
  M_composite = sum_i w_i M_i(T0, T)
Quantity/mix changes are reported separately and never labeled as inflation.

7. Uncertainty
Part-cluster bootstrap (default 100 iterations, seed 42) yields P10/P50/P90.
Intervals are widened for overall/category fallback, stale prices, and horizons > 24 months.

8. Limitations
- ~4 years of history; FY2026 incomplete.
- No supplier, currency, UoM, PO number, contract status, or facility fields.
- Apparent price changes may include supplier switches, spec changes, currency, contracts, or UoM changes.
- Most distinct parts are not multi-year; category/overall fallback is common.
- Horizons beyond 24 months are scenarios, not precise forecasts.
""".strip()


def write_results_workbook(path: Path, payload: dict[str, Any]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # --- Dashboard ---
    dash = wb.active
    dash.title = "Dashboard"
    dash["A1"] = "Parts Inflation Prototype — Dashboard"
    dash["A1"].font = Font(bold=True, size=16, color="1F4E79")
    meta = payload["dashboard"]
    labels = [
        ("Selected scope", meta.get("scope")),
        ("Data period", meta.get("data_period")),
        ("Base date", meta.get("base_date")),
        ("Target date", meta.get("target_date")),
        ("Composite P10 multiplier", meta.get("composite_p10")),
        ("Composite P50 / base multiplier", meta.get("composite_p50")),
        ("Composite P90 multiplier", meta.get("composite_p90")),
        ("Annualized implied inflation (P50)", meta.get("annualized_p50")),
        ("Projected fixed-basket cost change (P50)", meta.get("cost_change_p50")),
        ("Matched-spend coverage", meta.get("matched_spend_coverage")),
        ("% weight part-level", meta.get("pct_weight_part")),
        ("% weight category-level", meta.get("pct_weight_category")),
        ("% weight overall-level", meta.get("pct_weight_overall")),
        ("Selected model", meta.get("selected_model")),
        ("Selected forecast method", meta.get("forecast_method")),
        ("Long-horizon warning", meta.get("long_horizon_warning")),
    ]
    dash["A3"] = "Metric"
    dash["B3"] = "Value"
    _style_header(dash, 2)
    for i, (k, v) in enumerate(labels, 4):
        dash.cell(i, 1, k)
        cell = dash.cell(i, 2, v)
        if "warning" in k.lower() and v and str(v) not in {"", "None", "False"}:
            cell.fill = WARN_FILL
    # Historical index mini table for chart
    hist = payload.get("historical_index_chart")
    start = 4 + len(labels) + 2
    dash.cell(start, 1, "Historical / forecast index (chart data)")
    dash.cell(start, 1).font = Font(bold=True)
    if hist is not None and not hist.empty:
        for r_idx, row in enumerate(dataframe_to_rows(hist, index=False, header=True), start + 1):
            for c_idx, value in enumerate(row, 1):
                dash.cell(r_idx, c_idx, value)
        chart = LineChart()
        chart.title = "Historical index and forecast fan"
        chart.style = 10
        chart.y_axis.title = "Index (base≈1)"
        chart.x_axis.title = "Period"
        data = Reference(dash, min_col=2, min_row=start + 1, max_col=min(5, hist.shape[1] + 1), max_row=start + 1 + len(hist))
        cats = Reference(dash, min_col=1, min_row=start + 2, max_row=start + 1 + len(hist))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.shape = 4
        chart.width = 18
        chart.height = 10
        dash.add_chart(chart, "D4")

    cat_cmp = payload.get("category_comparison_chart")
    if cat_cmp is not None and not cat_cmp.empty:
        cstart = start + 3 + len(hist) if hist is not None and not hist.empty else start + 3
        dash.cell(cstart, 1, "Category forecast multipliers (P50)")
        dash.cell(cstart, 1).font = Font(bold=True)
        for r_idx, row in enumerate(dataframe_to_rows(cat_cmp, index=False, header=True), cstart + 1):
            for c_idx, value in enumerate(row, 1):
                dash.cell(r_idx, c_idx, value)
        chart2 = LineChart()
        chart2.title = "Category P50 multipliers"
        chart2.y_axis.title = "Multiplier"
        chart2.x_axis.title = "Category"
        # Use bar-like line; still informative
        data2 = Reference(dash, min_col=2, min_row=cstart + 1, max_col=2, max_row=cstart + 1 + len(cat_cmp))
        cats2 = Reference(dash, min_col=1, min_row=cstart + 2, max_row=cstart + 1 + len(cat_cmp))
        chart2.add_data(data2, titles_from_data=True)
        chart2.set_categories(cats2)
        chart2.width = 16
        chart2.height = 8
        dash.add_chart(chart2, "D20")
    _autosize(dash)

    def add_sheet(name: str, df: pd.DataFrame) -> None:
        ws = wb.create_sheet(name[:31])
        _write_df(ws, df if df is not None else pd.DataFrame())

    add_sheet("Controls Used", payload.get("controls_used"))
    add_sheet("Scope Sensitivity", payload.get("scope_sensitivity"))
    add_sheet("Historical Index", payload.get("historical_index"))
    add_sheet("Category Results", payload.get("category_results"))
    add_sheet("Part Forecasts", payload.get("part_forecasts"))
    add_sheet("Backtests", payload.get("backtests"))
    add_sheet("Scope Mapping", payload.get("scope_mapping"))
    add_sheet("Data Quality", payload.get("data_quality"))

    meth = wb.create_sheet("Methodology")
    meth["A1"] = "Methodology and Limitations"
    meth["A1"].font = Font(bold=True, size=14)
    meth["A3"] = METHODOLOGY_TEXT
    meth["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    meth.column_dimensions["A"].width = 120
    meth.row_dimensions[3].height = 420

    run = wb.create_sheet("Run Information")
    run_info = payload.get("run_information", pd.DataFrame())
    _write_df(run, run_info)

    # Profile comparison optional
    if payload.get("profile_comparison") is not None:
        add_sheet("Profile Check", payload["profile_comparison"])

    wb.save(path)
    logger.info("Wrote results workbook %s", path)
    return path
