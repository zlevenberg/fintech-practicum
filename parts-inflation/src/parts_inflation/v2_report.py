"""Client-facing and machine-readable outputs for the v2 pipeline."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def _json_value(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(v) for v in value]
    if isinstance(value, (pd.Timestamp, pd.Period)):
        return str(value)
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, float) and not np.isfinite(value):
        return None
    return value


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_table(ws, frame: pd.DataFrame) -> None:
    if frame is None or frame.empty:
        ws["A1"] = "No data available"
        return
    for row_number, row in enumerate(dataframe_to_rows(frame, index=False, header=True), 1):
        for column_number, value in enumerate(row, 1):
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            cell = ws.cell(row_number, column_number, value)
            if row_number == 1:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            elif column_number <= len(frame.columns):
                label = str(frame.columns[column_number - 1]).lower()
                if label == "cutoff" or label.endswith("_date") or label == "month_end":
                    cell.number_format = "yyyy-mm-dd"
                elif any(
                    token in label
                    for token in ["rate", "inflation", "escalation", "coverage", "share", "weight", "bias", "error"]
                ) and "pair_count" not in label:
                    cell.number_format = "0.00%"
                elif (
                    "spend" in label
                    or label.endswith("_value")
                    or (label == "value" and ws.title in {"Data Quality", "Exclusions"})
                ):
                    cell.number_format = "$#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = max((len(str(cell.value)) for cell in column[:200] if cell.value is not None), default=10)
        ws.column_dimensions[letter].width = min(max(width + 2, 11), 42)


def _methodology_frame() -> pd.DataFrame:
    rows = [
        ("Scope", "Combined direct costs: COS, Inventory, Operating Supplies, Production Supplies, Production Aids, and Small Tooling."),
        ("Realized price", "Extension (Qty Received) / Qty Received; Cost is only a documented fallback."),
        ("Committed price", "PO Value / Qty Ordered applied only to remaining open quantity; reported separately from realized inflation."),
        ("Repeat-sales observation", "y_j = ln(p_2/p_1); consecutive observed purchases may span multiple years."),
        ("Interval model", "y_j = sum_m D_jm*delta_m + beta*ln(q_2/q_1) + error, fit with Huber IRLS and smoothness/ridge penalties."),
        ("Identification", "The overall index and each cost bucket are fit separately; category-minus-overall paths are then shrunk toward zero."),
        ("Fixed basket", "M(T0,T) = sum_i q_i* p_i(T) / sum_i q_i* p_i(T0); latest complete FY realized spend is the default proxy."),
        ("Forecast", "Interpretable candidate time-series methods, bucket shrinkage, optional year-1 committed overlay, and mean reversion for years 2-3."),
        ("Backtest", "Rolling origins; actual price relatives are normalized to the exact 12-month horizon, gaps over 548 days are excluded as stale, and approved cutoff-only basket weights combine bucket results."),
        ("Uncertainty", "Comparison-entity cluster bootstrap; each successful iteration refits the selected estimator and forecast."),
        ("Fiscal year", "October 1 through September 30, labeled by ending year."),
        ("Mapping", "Exact and client-approved supersession matches are official. Heuristic family candidates are excluded until approved."),
    ]
    return pd.DataFrame(rows, columns=["topic", "method"])


def write_v2_outputs(run_dir: Path, payload: dict[str, Any]) -> dict[str, Path]:
    run_dir = Path(run_dir)
    run_dir.mkdir(parents=True, exist_ok=True)
    tables: dict[str, pd.DataFrame] = payload["tables"]
    csv_map = {
        "historical_index.csv": "historical_monthly",
        "fiscal_year_inflation.csv": "historical_rates",
        "forecast.csv": "forecast_all",
        "backtest_results.csv": "backtest_detail",
        "coverage.csv": "coverage",
        "data_quality.csv": "data_quality",
        "exclusion_log.csv": "exclusions",
        "pair_audit.csv": "pair_audit",
        "part_family_candidates.csv": "part_family_candidates",
        "committed_costs.csv": "committed_detail",
    }
    paths: dict[str, Path] = {}
    for filename, key in csv_map.items():
        path = run_dir / filename
        frame = tables.get(key, pd.DataFrame())
        frame.to_csv(path, index=False)
        paths[key] = path

    summary_path = run_dir / "summary.json"
    summary_path.write_text(
        json.dumps(_json_value(payload["summary"]), indent=2, sort_keys=True), encoding="utf-8"
    )
    paths["summary"] = summary_path

    config_path = run_dir / "resolved_config.json"
    config_path.write_text(
        json.dumps(_json_value(payload["resolved_config"]), indent=2, sort_keys=True), encoding="utf-8"
    )
    paths["resolved_config"] = config_path

    workbook_path = run_dir / "inflation_report.xlsx"
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Executive Summary"
    summary_ws["A1"] = "Glenair Direct-Cost Inflation — V2"
    summary_ws["A1"].font = Font(size=16, bold=True, color="1F4E79")
    summary_ws["A3"] = "Metric"
    summary_ws["B3"] = "Value"
    for cell in summary_ws[3]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    executive = payload["summary"].get("executive_metrics", {})
    for row_number, (key, value) in enumerate(executive.items(), 4):
        summary_ws.cell(row_number, 1, key)
        cell = summary_ws.cell(row_number, 2, _json_value(value))
        if "warning" in key.lower() or "limitation" in key.lower():
            cell.fill = WARN_FILL
        if isinstance(value, float) and any(token in key.lower() for token in ["rate", "inflation", "coverage", "escalation", "signal"]):
            cell.number_format = "0.00%"
    summary_ws.column_dimensions["A"].width = 42
    summary_ws.column_dimensions["B"].width = 85
    summary_ws.freeze_panes = "A4"
    summary_ws.sheet_view.showGridLines = False

    visual_ws = wb.create_sheet("Visual Dashboard", 1)
    visual_ws.sheet_view.showGridLines = False
    visual_ws.merge_cells("A1:Q1")
    visual_ws["A1"] = "Direct-Cost Inflation — Visual Dashboard"
    visual_ws["A1"].font = Font(size=16, bold=True, color="1F4E79")
    visual_ws.merge_cells("A2:Q2")
    visual_ws["A2"] = (
        "Backtest actuals are normalized to a 12-month horizon; stale base-to-actual gaps "
        "over 548 days are excluded, and category results use the approved cutoff-only basket."
    )
    visual_ws["A2"].font = Font(italic=True, color="666666")

    sheet_map = [
        ("Historical Fiscal Year", "historical_rates"),
        ("Historical TTM", "historical_ttm"),
        ("Historical Monthly", "historical_monthly"),
        ("Forecasts", "forecast_all"),
        ("Forecast Methods", "forecast_methods"),
        ("Forecast Monthly", "forecast_monthly"),
        ("Committed Costs", "committed_summary"),
        ("Scope & Coverage", "coverage"),
        ("Backtests", "backtest_summary"),
        ("Backtest Detail", "backtest_detail"),
        ("Data Quality", "data_quality"),
        ("Exclusions", "exclusions"),
        ("Mapping Candidates", "part_family_candidates"),
        ("Benchmark Monthly", "benchmark_monthly"),
        ("Benchmark Quarterly", "benchmark_quarterly"),
        ("Methodology", None),
        ("Run Information", "run_information"),
    ]
    for name, key in sheet_map:
        ws = wb.create_sheet(name[:31])
        _write_table(ws, _methodology_frame() if key is None else tables.get(key, pd.DataFrame()))

    monthly = tables.get("historical_monthly", pd.DataFrame())
    combined = monthly.loc[monthly.get("series", pd.Series(dtype=str)).eq("Combined Direct Costs")]
    if not combined.empty:
        chart_start = 40
        visual_ws.cell(chart_start, 1, "Month")
        visual_ws.cell(chart_start, 2, "Combined Direct-Cost Index")
        for offset, (_, row) in enumerate(combined.sort_values("month_end").iterrows(), 1):
            visual_ws.cell(chart_start + offset, 1, pd.Timestamp(row["month_end"]).strftime("%b %Y"))
            visual_ws.cell(chart_start + offset, 2, float(row["index"]))
        chart = LineChart()
        chart.title = "Combined Direct-Cost Index"
        chart.y_axis.title = "Index"
        chart.x_axis.title = "Month"
        data = Reference(
            visual_ws, min_col=2, min_row=chart_start,
            max_row=chart_start + len(combined),
        )
        cats = Reference(
            visual_ws, min_col=1, min_row=chart_start + 1,
            max_row=chart_start + len(combined),
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.visible_cells_only = False
        chart.legend = None
        chart.width = 18
        chart.height = 9
        summary_ws.add_chart(chart, "D3")

    forecast = tables.get("forecast_all", pd.DataFrame())
    composite = (
        forecast.loc[forecast["series"].eq("Combined Direct Costs")].sort_values("forecast_year")
        if {"series", "forecast_year"}.issubset(forecast.columns)
        else pd.DataFrame()
    )
    if not composite.empty:
        start = 40
        visual_ws.cell(start, 4, "Horizon")
        visual_ws.cell(start, 5, "Point")
        visual_ws.cell(start, 6, "Lower")
        visual_ws.cell(start, 7, "Upper")
        for offset, (_, row) in enumerate(composite.iterrows(), 1):
            visual_ws.cell(start + offset, 4, f"Year {int(row['forecast_year'])}")
            visual_ws.cell(start + offset, 5, float(row["cumulative_escalation"]))
            visual_ws.cell(start + offset, 6, float(row["lower"] - 1.0))
            visual_ws.cell(start + offset, 7, float(row["upper"] - 1.0))
        visual_ws[f"E{start + 1}:G{start + len(composite)}"][0][0].number_format = "0.00%"
        fc_chart = LineChart()
        fc_chart.title = "Projected cumulative escalation and bootstrap range"
        fc_chart.y_axis.title = "Cumulative escalation"
        fc_chart.y_axis.numFmt = "0%"
        fc_chart.add_data(
            Reference(visual_ws, min_col=5, max_col=7, min_row=start, max_row=start + len(composite)),
            titles_from_data=True,
        )
        fc_chart.set_categories(
            Reference(visual_ws, min_col=4, min_row=start + 1, max_row=start + len(composite))
        )
        for series, title in zip(fc_chart.series, ["Point forecast", "Lower bound", "Upper bound"]):
            series.tx = SeriesLabel(v=title)
        fc_chart.visible_cells_only = False
        fc_chart.width = 14
        fc_chart.height = 8
        visual_ws.add_chart(fc_chart, "A4")

    backtests = tables.get("backtest_detail", pd.DataFrame())
    selected_method = payload.get("manifest", {}).get("selected_forecast_method")
    selected_bt = (
        backtests.loc[backtests["method"].eq(selected_method)].sort_values("cutoff")
        if {"method", "cutoff"}.issubset(backtests.columns)
        else pd.DataFrame()
    )
    if not selected_bt.empty:
        start = 40
        visual_ws.cell(start, 9, "Cutoff")
        visual_ws.cell(start, 10, "Annualized actual")
        visual_ws.cell(start, 11, "Selected forecast")
        for offset, (_, row) in enumerate(selected_bt.iterrows(), 1):
            visual_ws.cell(start + offset, 9, pd.Timestamp(row["cutoff"]).strftime("%b %Y"))
            visual_ws.cell(start + offset, 10, float(row["actual_multiplier"] - 1.0))
            visual_ws.cell(start + offset, 11, float(row["predicted_multiplier"] - 1.0))
        bt_chart = LineChart()
        bt_chart.title = "Backtest: actual versus selected forecast"
        bt_chart.y_axis.title = "12-month inflation"
        bt_chart.y_axis.numFmt = "0%"
        bt_chart.add_data(
            Reference(visual_ws, min_col=10, max_col=11, min_row=start, max_row=start + len(selected_bt)),
            titles_from_data=True,
        )
        bt_chart.set_categories(
            Reference(visual_ws, min_col=9, min_row=start + 1, max_row=start + len(selected_bt))
        )
        for series, title in zip(bt_chart.series, ["Annualized actual", "Selected forecast"]):
            series.tx = SeriesLabel(v=title)
        bt_chart.visible_cells_only = False
        bt_chart.width = 14
        bt_chart.height = 8
        visual_ws.add_chart(bt_chart, "J4")

    coverage = tables.get("coverage", pd.DataFrame())
    bucket_forecast = (
        forecast.loc[
            forecast["forecast_year"].eq(1)
            & ~forecast["series"].eq("Combined Direct Costs")
        ].copy()
        if {"forecast_year", "series"}.issubset(forecast.columns)
        else pd.DataFrame()
    )
    if not bucket_forecast.empty and not coverage.empty:
        weight_map = coverage.set_index("category")["weight"]
        bucket_forecast["basket_weight"] = bucket_forecast["category"].map(weight_map)
        bucket_forecast["contribution"] = (
            bucket_forecast["basket_weight"] * bucket_forecast["annual_rate"]
        )
        bucket_forecast = bucket_forecast.sort_values("contribution")
        start = 49
        visual_ws.cell(start, 4, "Category")
        visual_ws.cell(start, 5, "Year 1 contribution")
        for offset, (_, row) in enumerate(bucket_forecast.iterrows(), 1):
            visual_ws.cell(start + offset, 4, str(row["category"]))
            visual_ws.cell(start + offset, 5, float(row["contribution"] * 10_000.0))
        driver_chart = BarChart()
        driver_chart.type = "bar"
        driver_chart.title = "Year 1 contribution by category (basis points)"
        driver_chart.x_axis.title = "Basis-point contribution"
        driver_chart.x_axis.numFmt = "0"
        driver_chart.add_data(
            Reference(visual_ws, min_col=5, min_row=start, max_row=start + len(bucket_forecast)),
            titles_from_data=True,
        )
        driver_chart.set_categories(
            Reference(visual_ws, min_col=4, min_row=start + 1, max_row=start + len(bucket_forecast))
        )
        driver_chart.visible_cells_only = False
        driver_chart.legend = None
        driver_chart.width = 14
        driver_chart.height = 8
        visual_ws.add_chart(driver_chart, "A21")

        coverage_plot = coverage[["category", "matched_spend_coverage"]].dropna().sort_values("matched_spend_coverage")
        start = 49
        visual_ws.cell(start, 9, "Category")
        visual_ws.cell(start, 10, "Matched-spend coverage")
        for offset, (_, row) in enumerate(coverage_plot.iterrows(), 1):
            visual_ws.cell(start + offset, 9, str(row["category"]))
            visual_ws.cell(start + offset, 10, float(row["matched_spend_coverage"] * 100.0))
        coverage_chart = BarChart()
        coverage_chart.type = "bar"
        coverage_chart.title = "Matched-spend coverage by category"
        coverage_chart.x_axis.title = "Coverage"
        coverage_chart.x_axis.numFmt = "0"
        coverage_chart.x_axis.scaling.min = 0
        coverage_chart.x_axis.scaling.max = 100
        coverage_chart.add_data(
            Reference(visual_ws, min_col=10, min_row=start, max_row=start + len(coverage_plot)),
            titles_from_data=True,
        )
        coverage_chart.set_categories(
            Reference(visual_ws, min_col=9, min_row=start + 1, max_row=start + len(coverage_plot))
        )
        coverage_chart.visible_cells_only = False
        coverage_chart.legend = None
        coverage_chart.width = 14
        coverage_chart.height = 8
        visual_ws.add_chart(coverage_chart, "J21")

    for col in range(1, 18):
        visual_ws.column_dimensions[get_column_letter(col)].width = 12
    for row in range(40, visual_ws.max_row + 1):
        visual_ws.row_dimensions[row].hidden = True
    wb.save(workbook_path)
    # Verify that the workbook can be reopened and all required sheets survived.
    check = load_workbook(workbook_path, read_only=True, data_only=False)
    required = {name[:31] for name, _ in sheet_map} | {
        "Executive Summary",
        "Visual Dashboard",
    }
    missing = required - set(check.sheetnames)
    check.close()
    if missing:
        raise RuntimeError(f"Excel output is missing sheets: {sorted(missing)}")
    paths["excel"] = workbook_path

    manifest_path = run_dir / "run_manifest.json"
    output_hashes = {path.name: _sha256(path) for path in paths.values() if path.exists()}
    manifest = dict(payload["manifest"])
    manifest["output_hashes"] = output_hashes
    manifest_path.write_text(json.dumps(_json_value(manifest), indent=2, sort_keys=True), encoding="utf-8")
    paths["manifest"] = manifest_path
    return paths
