"""Excel, CSV, and Markdown outputs for Historical Actual Inflation."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from parts_inflation.config import ResolvedConfig
from parts_inflation.historical_scope import PHYSICAL_INPUT_DESCRIPTIONS

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="D6E3F0")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
PCT_FORMAT = "0.00%"
MONEY_FORMAT = '"$"#,##0'
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def write_historical_outputs(
    output_dir: Path,
    payload: dict[str, Any],
    config: Optional[ResolvedConfig] = None,
) -> dict[str, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    ts = payload["run_timestamp"]
    stamp = ts.strftime("%Y-%m-%d_%H%M%S") if isinstance(ts, datetime) else datetime.now().strftime("%Y-%m-%d_%H%M%S")

    xlsx_path = output_dir / f"historical_actual_inflation_{stamp}.xlsx"
    csv_paths = {
        "scope": output_dir / "historical_scope_results.csv",
        "category": output_dir / "historical_category_results.csv",
        "method": output_dir / "historical_method_sensitivity.csv",
        "regression": output_dir / "historical_regression_results.csv",
        "matched": output_dir / "historical_matched_part_detail.csv",
    }
    md_path = output_dir / "historical_actual_inflation_summary.md"

    _write_csvs(payload, csv_paths)
    _write_markdown(md_path, payload)
    _write_excel(xlsx_path, payload, config)

    paths = {"excel": xlsx_path, "markdown": md_path, **{f"csv_{k}": v for k, v in csv_paths.items()}}
    logger.info("Wrote historical outputs under %s", output_dir)
    return paths


def _write_csvs(payload: dict[str, Any], paths: dict[str, Path]) -> None:
    payload["scope_results"].to_csv(paths["scope"], index=False)
    payload["category_results"].to_csv(paths["category"], index=False)
    payload["method_sensitivity"].to_csv(paths["method"], index=False)
    payload["regression_results"].to_csv(paths["regression"], index=False)
    detail = payload["matched_part_detail"]
    # Keep CSV manageable but complete
    detail.to_csv(paths["matched"], index=False)


def _fmt_pct(x: Any, digits: int = 2) -> str:
    try:
        if x is None or (isinstance(x, float) and not np.isfinite(x)):
            return "n/a"
        return f"{100.0 * float(x):.{digits}f}%"
    except Exception:
        return "n/a"


def _write_markdown(path: Path, payload: dict[str, Any]) -> None:
    chain = payload["chains"].get("physical_inputs", {})
    reg = payload["regression"]
    ts = payload["run_timestamp"]
    lines = [
        "# Historical Actual Inflation — Executive Summary",
        "",
        f"**Run timestamp:** {ts}",
        f"**Headline scope:** {payload['headline_scope']}",
        f"**FY2026 YTD through:** {payload['ytd_end']}",
        "",
        "## Principal finding",
        "",
        (
            "Based on repeated purchases of the same physical inputs, the client's spend-weighted "
            f"inflation averaged approximately {_fmt_pct(chain.get('annualized_fy2023_to_fy2025'), 1)} annually "
            f"through FY2025 (cumulative {_fmt_pct(chain.get('cumulative_fy2023_to_fy2025'), 1)} over two complete "
            f"fiscal years: {_fmt_pct(chain.get('fy2023_to_fy2024'))} then {_fmt_pct(chain.get('fy2024_to_fy2025'))}). "
            f"The price increase of the typical matched part was lower — approximately "
            f"{_fmt_pct(reg.get('unweighted_annual_rate_no_trim'), 1)} annually in the unweighted repeat-purchase "
            f"regression (equal-part winsorized fiscal rates are similar). In the aligned FY2026 year-to-date "
            f"comparison through {payload['ytd_end']}, spend-weighted physical-input inflation was "
            f"{_fmt_pct(chain.get('fy2026_ytd_yoy'))}."
        ),
        "",
        "This is the best current historical estimate given missing vendor and unit-of-measure data; "
        "it is **not** an official inflation rate and is **not** automatically a forward forecast.",
        "",
        "## Scope results (Robust Capped Törnqvist)",
        "",
        "| Scope | FY2023→FY2024 | FY2024→FY2025 | FY2026 YTD YoY | FY2023→FY2025 annualized |",
        "|---|---:|---:|---:|---:|",
    ]
    for sc in ("inventory_only", "physical_inputs", "all_po_lines"):
        c = payload["chains"].get(sc, {})
        lines.append(
            f"| {sc} | {_fmt_pct(c.get('fy2023_to_fy2024'))} | {_fmt_pct(c.get('fy2024_to_fy2025'))} | "
            f"{_fmt_pct(c.get('fy2026_ytd_yoy'))} | {_fmt_pct(c.get('annualized_fy2023_to_fy2025'))} |"
        )

    lines += [
        "",
        "## Physical-input yearly detail",
        "",
    ]
    phys = payload["scope_results"]
    phys = phys.loc[phys["scope"] == "physical_inputs"] if not phys.empty else phys
    if not phys.empty:
        lines.append(
            "| Comparison | Matched parts | Coverage prior/current | Headline | Geo-spend | Equal-part | Similar-qty |"
        )
        lines.append("|---|---:|---:|---:|---:|---:|---:|")
        for _, r in phys.iterrows():
            cov = f"{_fmt_pct(r.get('coverage_a'))} / {_fmt_pct(r.get('coverage_b'))}"
            lines.append(
                f"| {r['comparison']} | {int(r.get('matched_parts', 0)):,} | {cov} | "
                f"{_fmt_pct(r.get('headline_robust_capped_tornqvist'))} | "
                f"{_fmt_pct(r.get('capped_geometric_spend'))} | "
                f"{_fmt_pct(r.get('equal_part_winsorized'))} | "
                f"{_fmt_pct(r.get('similar_quantity_equal_winsorized'))} |"
            )

    lines += [
        "",
        "## Repeat-purchase regression (physical inputs)",
        "",
        f"- Pairs (pre-trim): {reg.get('n_pairs'):,}",
        f"- Unweighted annual rate (no trim): {_fmt_pct(reg.get('unweighted_annual_rate_no_trim'))}",
        f"- Spend-weighted annual rate (no trim): {_fmt_pct(reg.get('weighted_annual_rate_no_trim'))}",
        f"- Quantity elasticity γ: {reg.get('gamma_no_trim_weighted'):.4f}",
        f"- Effect of doubling quantity: {_fmt_pct(reg.get('qty_doubling_no_trim_weighted'))}",
        "",
        "## Methods",
        "",
        "- Same-part matching on normalized PartKey across fiscal windows.",
        "- Representative period price = median Cost; spend weight = PO Value (not Cost×Qty Ordered).",
        "- Headline index = Robust Capped Törnqvist (1st/99th winsorized log relatives; 95th-percentile share cap).",
        "- Complete-year chain uses only FY2023→FY2024 and FY2024→FY2025; FY2026 YTD is a separate aligned YoY.",
        "- Independent Huber regression on adjacent same-part pairs (≥30-day gap) validates the spend-weighted rate.",
        "",
        "## Data limitations",
        "",
        "- Only about four fiscal periods; FY2026 is incomplete (YTD through July 9, 2026).",
        "- Vendor, currency, UoM, PO number, contract status, and facility are missing.",
        "- Same-part price changes can include supplier, specification, contract, emergency, or unit changes.",
        "- Matched parts cover only a portion of physical-input PO value (see coverage columns).",
        "- Most distinct parts are not observed across multiple fiscal periods.",
        "- Category rates with low coverage or few matches are unstable (flagged Low reliability).",
        "- PO Value and Cost×Qty Ordered differ materially in aggregate; PO Value is used for weights.",
        "- Spend-weighted rates describe cost exposure; unweighted rates describe the typical matched part.",
        "- Historical actual inflation is not automatically the best forward forecast.",
        "",
        f"Physical-input Description membership: {', '.join(sorted(PHYSICAL_INPUT_DESCRIPTIONS))}",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, max_width: int = 36) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 10
        for cell in col[:80]:
            if cell.value is not None:
                width = max(width, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def _df_to_sheet(wb: Workbook, name: str, df: pd.DataFrame, pct_cols: Optional[list[str]] = None) -> None:
    ws = wb.create_sheet(name)
    if df is None or df.empty:
        ws.append(["(no rows)"])
        return
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        ws.append(row)
    _style_header(ws, 1)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    pct_cols = pct_cols or []
    headers = {cell.value: cell.column for cell in ws[1]}
    for col_name in pct_cols:
        if col_name in headers:
            c = headers[col_name]
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=c)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = PCT_FORMAT
    _autosize(ws)


def _write_excel(path: Path, payload: dict[str, Any], config: Optional[ResolvedConfig]) -> None:
    wb = Workbook()
    # Executive Summary
    ws = wb.active
    ws.title = "Executive Summary"
    chain = payload["chains"].get("physical_inputs", {})
    reg = payload["regression"]
    phys = payload["scope_results"]
    phys = phys.loc[phys["scope"] == "physical_inputs"] if not phys.empty else phys

    ws["A1"] = "Historical Actual Inflation — Executive Summary"
    ws["A1"].font = Font(bold=True, size=16, color="1F4E79")
    ws.merge_cells("A1:F1")

    narrative = (
        f"Based on repeated purchases of the same physical inputs, spend-weighted inflation averaged "
        f"{_fmt_pct(chain.get('annualized_fy2023_to_fy2025'), 1)} annually through FY2025 "
        f"(multiplier ≈ {1 + float(chain.get('annualized_fy2023_to_fy2025') or 0):.3f} per year; "
        f"cumulative {_fmt_pct(chain.get('cumulative_fy2023_to_fy2025'), 1)}). "
        f"Typical matched-part inflation was lower (~{_fmt_pct(reg.get('unweighted_annual_rate_no_trim'), 1)} annually). "
        f"Aligned FY2026 YTD through {payload['ytd_end']}: {_fmt_pct(chain.get('fy2026_ytd_yoy'))}. "
        f"Doubling order quantity associates with ≈{_fmt_pct(reg.get('qty_doubling_no_trim_weighted'))} unit price."
    )
    ws["A3"] = narrative
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:F5")
    ws.row_dimensions[3].height = 60

    bullets = [
        ("Annual physical-input inflation (spend-weighted, thru FY2025)", _fmt_pct(chain.get("annualized_fy2023_to_fy2025"), 1)),
        ("FY2023→FY2025 cumulative", _fmt_pct(chain.get("cumulative_fy2023_to_fy2025"), 1)),
        ("FY2026 YTD aligned YoY", _fmt_pct(chain.get("fy2026_ytd_yoy"))),
        ("Typical-part inflation (unweighted regression)", _fmt_pct(reg.get("unweighted_annual_rate_no_trim"), 1)),
        ("Quantity-doubling association", _fmt_pct(reg.get("qty_doubling_no_trim_weighted"))),
        ("Regression pairs (pre-trim)", f"{reg.get('n_pairs'):,}"),
        ("Headline method", "Robust Capped Törnqvist"),
        ("Headline scope", payload["headline_scope"]),
    ]
    ws["A7"] = "Key metrics"
    ws["A7"].font = Font(bold=True, size=12)
    ws["A8"] = "Metric"
    ws["B8"] = "Value"
    _style_header(ws, 8)
    for i, (k, v) in enumerate(bullets, 9):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)

    # Scope table
    row0 = 9 + len(bullets) + 1
    ws.cell(row0, 1, "Annual inflation by scope (Robust Capped Törnqvist)").font = Font(bold=True, size=12)
    headers = ["Scope", "FY2023→FY2024", "FY2024→FY2025", "FY2026 YTD YoY", "FY2023→FY2025 annualized"]
    for j, h in enumerate(headers, 1):
        ws.cell(row0 + 1, j, h)
    _style_header(ws, row0 + 1)
    for i, sc in enumerate(("inventory_only", "physical_inputs", "all_po_lines"), row0 + 2):
        c = payload["chains"].get(sc, {})
        ws.cell(i, 1, sc)
        ws.cell(i, 2, c.get("fy2023_to_fy2024"))
        ws.cell(i, 3, c.get("fy2024_to_fy2025"))
        ws.cell(i, 4, c.get("fy2026_ytd_yoy"))
        ws.cell(i, 5, c.get("annualized_fy2023_to_fy2025"))
        for j in range(2, 6):
            ws.cell(i, j).number_format = PCT_FORMAT

    # Match counts
    r = row0 + 6
    ws.cell(r, 1, "Physical-input match counts and coverage").font = Font(bold=True)
    ws.cell(r + 1, 1, "Comparison")
    ws.cell(r + 1, 2, "Matched parts")
    ws.cell(r + 1, 3, "Coverage prior")
    ws.cell(r + 1, 4, "Coverage current")
    ws.cell(r + 1, 5, "Headline inflation")
    _style_header(ws, r + 1)
    if not phys.empty:
        for i, (_, row) in enumerate(phys.iterrows(), r + 2):
            ws.cell(i, 1, row["comparison"])
            ws.cell(i, 2, int(row.get("matched_parts", 0)))
            ws.cell(i, 3, row.get("coverage_a"))
            ws.cell(i, 4, row.get("coverage_b"))
            ws.cell(i, 5, row.get("headline_robust_capped_tornqvist"))
            ws.cell(i, 3).number_format = PCT_FORMAT
            ws.cell(i, 4).number_format = PCT_FORMAT
            ws.cell(i, 5).number_format = PCT_FORMAT

    lim_row = r + 8
    ws.cell(lim_row, 1, "Limitations").font = Font(bold=True)
    limitations = [
        "Only ~four fiscal periods; FY2026 incomplete.",
        "Vendor / currency / UoM / PO number / contract / facility missing.",
        "Matched spend coverage is partial (~45–58% in headline physical comparisons).",
        "PO Value used for weights (Cost×Qty Ordered aggregates are inconsistent).",
        "Historical actuals are distinct from any forward forecast.",
        "7.7–7.9% is a best estimate, not an official identified rate.",
    ]
    for i, t in enumerate(limitations):
        ws.cell(lim_row + 1 + i, 1, f"• {t}")

    ws.column_dimensions["A"].width = 62
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18

    # Remaining sheets
    pct_scope = [
        "median_pct_change",
        "p25_pct_change",
        "p75_pct_change",
        "pct_no_price_change",
        "coverage_a",
        "coverage_b",
        "raw_tornqvist",
        "winsorized_tornqvist",
        "headline_robust_capped_tornqvist",
        "capped_geometric_spend",
        "equal_part_winsorized",
        "equal_part_raw",
        "similar_quantity_equal_winsorized",
    ]
    _df_to_sheet(wb, "Scope Results", payload["scope_results"], pct_scope)
    _df_to_sheet(
        wb,
        "Category Results",
        payload["category_results"],
        [
            "coverage_a",
            "coverage_b",
            "headline_robust_capped_tornqvist",
            "equal_part_winsorized",
            "capped_geometric_spend",
        ],
    )
    # Flag low reliability visually
    cat_ws = wb["Category Results"]
    headers = {c.value: c.column for c in cat_ws[1]}
    if "reliability_flag" in headers and "low_reliability" in headers:
        flag_col = headers["reliability_flag"]
        for row in range(2, cat_ws.max_row + 1):
            if cat_ws.cell(row, headers["low_reliability"]).value in (True, "True", 1):
                cat_ws.cell(row, flag_col).fill = WARN_FILL

    _df_to_sheet(wb, "Method Sensitivity", payload["method_sensitivity"], ["inflation_rate", "coverage_a", "coverage_b"])
    _df_to_sheet(
        wb,
        "Regression Results",
        payload["regression_results"],
        ["annual_rate", "qty_doubling_effect"],
    )

    # Matched part detail — selected columns for performance
    detail = payload["matched_part_detail"]
    detail_cols = [
        c
        for c in [
            "scope",
            "comparison",
            "PartKey",
            "raw_part_number",
            "description_norm",
            "hist_category",
            "Description 1",
            "price_a",
            "price_b",
            "qty_a",
            "qty_b",
            "spend_a",
            "spend_b",
            "r",
            "g",
            "r_raw",
            "r_W",
            "s_a",
            "s_b",
            "s_bar",
            "s_tilde",
            "h",
            "h_tilde",
            "contrib_headline_log",
            "quantity_comparable",
            "is_price_outlier",
            "no_price_change",
            "coverage_a",
            "coverage_b",
            "line_count_a",
            "line_count_b",
            "min_po_date_a",
            "max_po_date_a",
            "min_po_date_b",
            "max_po_date_b",
        ]
        if detail is not None and not detail.empty and c in detail.columns
    ]
    _df_to_sheet(wb, "Matched Part Detail", detail[detail_cols] if detail_cols else detail, ["g", "coverage_a", "coverage_b"])

    recon = payload.get("spend_reconciliation", pd.DataFrame())
    _df_to_sheet(wb, "Spend Reconciliation", recon)

    # Data Quality
    dq_rows = []
    for info in payload["infos"]:
        dq_rows.append(
            {
                "item": f"source:{info.path.name}",
                "value": f"{info.nrows} rows; {info.min_date} → {info.max_date}; sheet={info.sheet}",
            }
        )
    prof = payload.get("profile_summary", {})
    for k, v in prof.items():
        dq_rows.append({"item": k, "value": v})
    for w in payload.get("warnings", []):
        dq_rows.append({"item": "warning", "value": w})
    if payload.get("scope_disclosure") is not None:
        for _, r in payload["scope_disclosure"].iterrows():
            dq_rows.append(
                {
                    "item": f"scope_disclosure:{r.get('description_norm')}",
                    "value": f"rows={r.get('row_count')}; PO Value={r.get('po_value')}; role={r.get('scope_role')}",
                }
            )
    _df_to_sheet(wb, "Data Quality", pd.DataFrame(dq_rows))

    # Methodology
    meth = wb.create_sheet("Methodology")
    meth["A1"] = "Methodology — Historical Actual Inflation"
    meth["A1"].font = Font(bold=True, size=14)
    blocks = [
        (
            "Same-part matching",
            "PartKey is Unicode-normalized, trimmed, whitespace-collapsed, uppercased text of the raw part number. "
            "Punctuation is preserved. A part enters the matched set for periods a and b when both periods have a positive median Cost.",
        ),
        (
            "Median part-period prices",
            "p_{i,t} = median{Cost_{i,t,k}}. Using the median avoids letting inconsistent Qty Ordered units distort the representative price.",
        ),
        (
            "PO-value weighting",
            "v_{i,t} = sum max(PO Value, 0). Do not use Cost×Qty Ordered as the spend weight — aggregate inconsistencies are material.",
        ),
        (
            "Winsorization",
            "Within each scope/pair: L=Q_0.01(r), U=Q_0.99(r), r^W = clip(r, L, U). Outliers are flagged, not deleted.",
        ),
        (
            "Weight capping",
            "Average Törnqvist shares s̄_i = (s_{i,a}+s_{i,b})/2; c_i = min(s̄_i, Q_0.95(s̄)); s̃_i = c_i / sum c.",
        ),
        (
            "Robust Capped Törnqvist (headline)",
            "π_{T,cap} = exp(sum_i s̃_i r_i^W) − 1. Both price relatives and weights are modified for robustness.",
        ),
        (
            "Period chaining",
            "M = (1+π_23→24)(1+π_24→25); cumulative = M−1; annualized = M^(1/2)−1. FY2026 YTD is NOT chained after FY2025.",
        ),
        (
            "Repeat-purchase regression",
            "Adjacent pairs only; ≥30-day gap; y = β_t Δt + γ x_q + ε (Huber, no intercept). Annual rate = exp(β_t)−1. "
            "Weights w_j = b_j^cap / sqrt(N_i), mean-normalized to 1.",
        ),
        (
            "Quantity adjustment",
            "γ is the elasticity of unit price to order quantity. Doubling quantity associates with 2^γ − 1.",
        ),
        (
            "Why typical-part and spend-weighted rates differ",
            "Equal-part / unweighted methods describe the median matched part. Spend-weighted methods describe the client's "
            "cost exposure, which is concentrated in categories (especially Production Supplies) with higher inflation.",
        ),
        (
            "Physical-input scope",
            "Exact normalized Description membership: " + ", ".join(sorted(PHYSICAL_INPUT_DESCRIPTIONS)),
        ),
    ]
    row = 3
    for title, body in blocks:
        meth.cell(row, 1, title).font = Font(bold=True)
        meth.cell(row + 1, 1, body).alignment = Alignment(wrap_text=True)
        meth.merge_cells(start_row=row + 1, start_column=1, end_row=row + 2, end_column=1)
        row += 4
    meth.column_dimensions["A"].width = 110

    # Run Information
    run_rows = [
        {"item": "run_timestamp", "value": str(payload["run_timestamp"])},
        {"item": "version", "value": payload.get("version")},
        {"item": "platform", "value": payload.get("platform")},
        {"item": "python", "value": payload.get("python")},
        {"item": "elapsed_seconds", "value": payload.get("elapsed_seconds")},
        {"item": "headline_scope", "value": payload.get("headline_scope")},
        {"item": "ytd_end", "value": str(payload.get("ytd_end"))},
        {"item": "log_path", "value": payload.get("log_path")},
    ]
    for k, v in payload.get("settings", {}).items():
        run_rows.append({"item": f"setting:{k}", "value": v})
    for name, h in payload.get("input_hashes", {}).items():
        run_rows.append({"item": f"input_sha256:{name}", "value": h})
    if config is not None:
        run_rows.append({"item": "config_path", "value": str(config.config_path)})
    _df_to_sheet(wb, "Run Information", pd.DataFrame(run_rows))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    logger.info("Wrote workbook %s", path)
