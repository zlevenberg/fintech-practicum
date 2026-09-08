"""Validated configuration loading, defaults, and Excel config workbook I/O."""

from __future__ import annotations

import logging
from dataclasses import dataclass, field, fields
from datetime import date, datetime, timedelta
from enum import Enum
from pathlib import Path
from typing import Any, Optional

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, field_validator

logger = logging.getLogger(__name__)

REQUIRED_SOURCE_COLUMNS = [
    "P.O. Date",
    "Part Number",
    "Description 1",
    "Description 2",
    "Cost",
    "Qty Ordered",
    "Qty Received",
    "Extension (Qty Received)",
    "PO Value",
    "Bucket",
    "Description",
]


class ScopeMode(str, Enum):
    inventory_only = "inventory_only"
    physical_inputs = "physical_inputs"
    all_po_lines = "all_po_lines"


class CompositeWeighting(str, Enum):
    trailing_12m_po_value = "trailing_12m_po_value"
    planned_basket = "planned_basket"


class SameDayAgg(str, Enum):
    quantity_weighted_mean = "quantity_weighted_mean"
    median = "median"


class QuantityAdjustmentMode(str, Enum):
    estimate = "estimate"
    ignore = "ignore"
    zero = "zero"


class DedupMode(str, Enum):
    flag_only = "flag_only"
    drop = "drop"


class SelectedModelMode(str, Enum):
    best_backtest = "best_backtest"
    hierarchical = "hierarchical"
    tornqvist = "tornqvist"
    matched_part = "matched_part"
    last_price = "last_price"
    overall_cagr = "overall_cagr"
    category_benchmark = "category_benchmark"
    blended = "blended"


class ControlDefaults(BaseModel):
    scope_mode: ScopeMode = ScopeMode.physical_inputs
    base_date: Optional[date] = None
    target_date: Optional[date] = None
    composite_weighting: CompositeWeighting = CompositeWeighting.trailing_12m_po_value
    price_field: str = "Cost"
    quantity_field: str = "Qty Ordered"
    spend_field: str = "PO Value"
    include_open_orders_as_prices: bool = True
    include_open_orders_in_weights: bool = True
    same_day_price_aggregation: SameDayAgg = SameDayAgg.quantity_weighted_mean
    quantity_adjustment_mode: QuantityAdjustmentMode = QuantityAdjustmentMode.estimate
    quantity_fallback_to_received: bool = False
    deduplication_mode: DedupMode = DedupMode.flag_only
    benchmark_winsor_lower: float = 0.01
    benchmark_winsor_upper: float = 0.99
    part_min_intervals: int = 3
    part_min_span_days: int = 365
    category_min_pairs: int = 100
    part_shrinkage_k: float = 5.0
    forecast_horizons_months: str = "3,6,12"
    bootstrap_iterations: int = 100
    random_seed: int = 42
    confidence_lower_quantile: float = 0.10
    confidence_upper_quantile: float = 0.90
    long_horizon_warning_months: int = 24
    selected_model_mode: SelectedModelMode = SelectedModelMode.best_backtest
    fast_mode: bool = False
    material_wape_improvement: float = 0.01
    extreme_ratio_low: float = 0.25
    extreme_ratio_high: float = 4.0
    extreme_ratio_max_days: int = 548  # ~18 months
    pair_weight_method: str = "capped_spend_sqrt_n"
    lambda_smooth: float = 10.0
    lambda_ridge: float = 1.0
    lambda_u: float = 5.0
    lambda_us: float = 5.0
    lambda_gamma: float = 2.0
    huber_delta: float = 1.5
    cache_enabled: bool = True

    @field_validator(
        "include_open_orders_as_prices",
        "include_open_orders_in_weights",
        "quantity_fallback_to_received",
        "fast_mode",
        "cache_enabled",
        mode="before",
    )
    @classmethod
    def parse_bool(cls, v: Any) -> bool:
        if isinstance(v, bool):
            return v
        if v is None:
            return False
        s = str(v).strip().upper()
        if s in {"TRUE", "1", "YES", "Y"}:
            return True
        if s in {"FALSE", "0", "NO", "N", ""}:
            return False
        raise ValueError(f"Invalid boolean value: {v!r}")

    @field_validator("base_date", "target_date", mode="before")
    @classmethod
    def parse_date(cls, v: Any) -> Optional[date]:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return None
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        if isinstance(v, str) and not v.strip():
            return None
        return pd.to_datetime(v).date()

    def horizon_months(self) -> list[int]:
        return [int(x.strip()) for x in self.forecast_horizons_months.split(",") if x.strip()]


CONTROL_DESCRIPTIONS: dict[str, str] = {
    "scope_mode": "Analysis scope: inventory_only | physical_inputs | all_po_lines",
    "base_date": "Model base date T0 (blank = latest valid PO date)",
    "target_date": "Forecast target date (blank = base_date + 12 months)",
    "composite_weighting": "Composite basket weights: trailing_12m_po_value | planned_basket",
    "price_field": "Unit price field (default Cost)",
    "quantity_field": "Quantity field (default Qty Ordered)",
    "spend_field": "Spend/weight field (default PO Value)",
    "include_open_orders_as_prices": "Treat zero-received positive ordered rows as price observations",
    "include_open_orders_in_weights": "Include open orders in ordered-spend weights",
    "same_day_price_aggregation": "quantity_weighted_mean | median",
    "quantity_adjustment_mode": "estimate | ignore | zero quantity elasticity",
    "quantity_fallback_to_received": "If ordered qty unusable, fall back to positive Qty Received",
    "deduplication_mode": "flag_only | drop exact duplicate fingerprints",
    "benchmark_winsor_lower": "Lower quantile for winsorized benchmark log changes",
    "benchmark_winsor_upper": "Upper quantile for winsorized benchmark log changes",
    "part_min_intervals": "Minimum adjacent pairs for part-specific residual",
    "part_min_span_days": "Minimum observation span (days) for part-specific residual",
    "category_min_pairs": "Minimum pairs before category deviation is trusted",
    "part_shrinkage_k": "Shrinkage prior strength k for part residuals",
    "forecast_horizons_months": "Comma-separated backtest horizons in months",
    "bootstrap_iterations": "Part-cluster bootstrap iterations for P10/P90",
    "random_seed": "RNG seed for reproducibility",
    "confidence_lower_quantile": "Lower confidence quantile (default 0.10)",
    "confidence_upper_quantile": "Upper confidence quantile (default 0.90)",
    "long_horizon_warning_months": "Horizons beyond this are labeled scenarios",
    "selected_model_mode": "best_backtest or force a specific model family",
    "fast_mode": "Reduce bootstrap iterations and lambda grid size only (does not subsample pairs)",
    "material_wape_improvement": "Relative WAPE improvement required to prefer complexity",
    "extreme_ratio_low": "Flag price ratios below this over short intervals",
    "extreme_ratio_high": "Flag price ratios above this over short intervals",
    "extreme_ratio_max_days": "Max interval days for extreme ratio flags (~18 months)",
    "pair_weight_method": "capped_spend_sqrt_n | equal_part",
    "lambda_smooth": "Smoothness penalty on monthly inflation first differences",
    "lambda_ridge": "Ridge penalty on overall monthly inflation",
    "lambda_u": "Ridge penalty on category monthly deviations",
    "lambda_us": "Smoothness penalty on category deviation differences",
    "lambda_gamma": "Ridge penalty on category quantity-elasticity deviations",
    "huber_delta": "Huber loss threshold in residual sigma units",
    "cache_enabled": "Use Parquet cache for cleaned data and pairs",
}


@dataclass
class SettingSource:
    value: Any
    source: str  # CLI | Config | Default
    description: str = ""


@dataclass
class ResolvedConfig:
    controls: ControlDefaults
    sources: dict[str, SettingSource] = field(default_factory=dict)
    scope_mapping: pd.DataFrame = field(default_factory=pd.DataFrame)
    category_mapping: pd.DataFrame = field(default_factory=pd.DataFrame)
    part_overrides: pd.DataFrame = field(default_factory=pd.DataFrame)
    planned_basket: pd.DataFrame = field(default_factory=pd.DataFrame)
    config_path: Optional[Path] = None

    def get(self, key: str) -> Any:
        return getattr(self.controls, key)


def _normalize_header(name: Any) -> str:
    if name is None:
        return ""
    s = str(name).strip()
    if s.lower() == "cost" or s.replace(" ", "").lower() == "cost":
        return "Cost"
    return s


def project_root() -> Path:
    return Path(__file__).resolve().parents[2]


# Heuristic include/exclude keywords for Scope Mapping
INCLUDE_KEYWORDS = [
    "inventory",
    "production supplies",
    "operating supplies",
    "production aids",
    "small tools",
    "office supplies",
    "machinery",
    "furnishings",
    "product testing",
]
EXCLUDE_KEYWORDS = [
    "outside services",
    "labor",
    "shipping",
    "freight",
    "consulting",
    "repairs & maint",
    "repairs and maint",
]
NEEDS_REVIEW_KEYWORDS = [
    "unsure",
    "needs review",
    "other",
    "misc",
]


def classify_scope_row(bucket: Any, description: Any) -> tuple[str, str, str]:
    """Return (default_decision, physical_input_category, reason)."""
    desc = "" if description is None or (isinstance(description, float) and pd.isna(description)) else str(description)
    bucket_s = "" if bucket is None or (isinstance(bucket, float) and pd.isna(bucket)) else str(bucket)
    text = f"{bucket_s} {desc}".lower()
    cat = desc.strip() if desc.strip() else "Uncategorized"

    for kw in EXCLUDE_KEYWORDS:
        if kw in text:
            return "Exclude", cat, f"Matched exclude keyword: {kw}"
    for kw in NEEDS_REVIEW_KEYWORDS:
        if kw in text:
            return "Needs Review", cat, f"Ambiguous category keyword: {kw}"
    for kw in INCLUDE_KEYWORDS:
        if kw in text:
            return "Include", cat, f"Matched include keyword: {kw}"
    return "Needs Review", cat, "No strong include/exclude keyword; marked Needs Review"


def invent_initial_scope_mapping(unique_pairs: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, r in unique_pairs.iterrows():
        decision, cat, reason = classify_scope_row(r.get("Bucket"), r.get("Description"))
        rows.append(
            {
                "Bucket": r.get("Bucket"),
                "Description": r.get("Description"),
                "Default Scope Decision": decision,
                "Physical Input Category": cat,
                "Reason": reason,
                "Client Approved": False,
                "Manual Override": "",
            }
        )
    return pd.DataFrame(rows)


def write_config_workbook(
    path: Path,
    scope_mapping: Optional[pd.DataFrame] = None,
    controls: Optional[ControlDefaults] = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    controls = controls or ControlDefaults()
    wb = openpyxl.Workbook()

    # Controls
    ws = wb.active
    ws.title = "Controls"
    header = ["Key", "Value", "Description"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E2F3")
    for key, desc in CONTROL_DESCRIPTIONS.items():
        val = getattr(controls, key)
        if isinstance(val, Enum):
            val = val.value
        elif isinstance(val, bool):
            val = "TRUE" if val else "FALSE"
        elif isinstance(val, date):
            val = val.isoformat()
        elif val is None:
            val = ""
        ws.append([key, val, desc])
    for col in range(1, 4):
        ws.column_dimensions[get_column_letter(col)].width = [28, 28, 70][col - 1]
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    def _write_sheet(name: str, columns: list[str], df: Optional[pd.DataFrame]) -> None:
        w = wb.create_sheet(name)
        w.append(columns)
        for cell in w[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E2F3")
        if df is not None and not df.empty:
            for _, row in df.iterrows():
                w.append([row.get(c, "") for c in columns])
        for i, c in enumerate(columns, 1):
            w.column_dimensions[get_column_letter(i)].width = min(40, max(14, len(c) + 2))
        w.auto_filter.ref = w.dimensions
        w.freeze_panes = "A2"

    scope_cols = [
        "Bucket",
        "Description",
        "Default Scope Decision",
        "Physical Input Category",
        "Reason",
        "Client Approved",
        "Manual Override",
    ]
    if scope_mapping is None:
        scope_mapping = pd.DataFrame(columns=scope_cols)
    _write_sheet("Scope Mapping", scope_cols, scope_mapping)

    _write_sheet(
        "Category Mapping",
        ["Bucket", "Description", "Physical Input Category", "Notes"],
        None,
    )
    _write_sheet(
        "Part Overrides",
        [
            "PartKey",
            "IncludeExclude",
            "ReplacementPartKey",
            "Category",
            "UoMAdjustmentFactor",
            "ManualCurrentPrice",
            "ManualFutureQuantity",
            "Notes",
        ],
        None,
    )
    _write_sheet(
        "Planned Basket",
        ["PartKey", "ExpectedQuantity", "ExpectedSpend", "EffectiveDate", "Notes"],
        None,
    )

    # Methodology notes sheet for auditors editing config
    notes = wb.create_sheet("Config Notes")
    notes["A1"] = "Config Notes"
    notes["A1"].font = Font(bold=True, size=14)
    notes["A3"] = (
        "Manual Override on Scope Mapping always wins over Default Scope Decision. "
        "Needs Review rows are excluded from physical_inputs by default but quantified in reports. "
        "Leave base_date / target_date blank to use latest PO date and +12 months."
    )
    notes["A3"].alignment = Alignment(wrap_text=True)
    notes.column_dimensions["A"].width = 100

    wb.save(path)
    logger.info("Wrote configuration workbook to %s", path)
    return path


def read_controls_sheet(path: Path) -> dict[str, Any]:
    df = pd.read_excel(path, sheet_name="Controls")
    df.columns = [str(c).strip() for c in df.columns]
    if "Key" not in df.columns or "Value" not in df.columns:
        raise ValueError(f"Controls sheet in {path} must have Key and Value columns")
    out: dict[str, Any] = {}
    for _, row in df.iterrows():
        key = str(row["Key"]).strip()
        if not key or key.lower() == "nan":
            continue
        out[key] = row["Value"]
    return out


def _read_optional_sheet(path: Path, name: str) -> pd.DataFrame:
    try:
        df = pd.read_excel(path, sheet_name=name)
        df.columns = [str(c).strip() for c in df.columns]
        return df
    except ValueError:
        return pd.DataFrame()


def load_config(
    path: Optional[Path] = None,
    cli_overrides: Optional[dict[str, Any]] = None,
    create_if_missing: bool = True,
    scope_seed: Optional[pd.DataFrame] = None,
) -> ResolvedConfig:
    """Load config with precedence: CLI > Config workbook > code defaults."""
    cli_overrides = {k: v for k, v in (cli_overrides or {}).items() if v is not None}
    root = project_root()
    path = Path(path) if path else root / "config" / "model_config.xlsx"

    if not path.exists():
        if not create_if_missing:
            raise FileNotFoundError(f"Config not found: {path}")
        write_config_workbook(path, scope_mapping=scope_seed)

    raw = read_controls_sheet(path)
    sources: dict[str, SettingSource] = {}
    merged: dict[str, Any] = {}

    defaults = ControlDefaults()
    for f in defaults.model_fields:
        default_val = getattr(defaults, f)
        if f in cli_overrides:
            merged[f] = cli_overrides[f]
            sources[f] = SettingSource(cli_overrides[f], "CLI", CONTROL_DESCRIPTIONS.get(f, ""))
        elif f in raw and not (isinstance(raw[f], float) and pd.isna(raw[f])) and str(raw[f]).strip() != "":
            merged[f] = raw[f]
            sources[f] = SettingSource(raw[f], "Config", CONTROL_DESCRIPTIONS.get(f, ""))
        else:
            merged[f] = default_val
            sources[f] = SettingSource(default_val, "Default", CONTROL_DESCRIPTIONS.get(f, ""))

    try:
        controls = ControlDefaults(**merged)
    except Exception as exc:
        raise ValueError(f"Invalid configuration values in {path}: {exc}") from exc

    if controls.fast_mode:
        if "bootstrap_iterations" not in cli_overrides:
            controls.bootstrap_iterations = min(controls.bootstrap_iterations, 10)
            sources["bootstrap_iterations"] = SettingSource(
                controls.bootstrap_iterations, "Default", "Reduced by fast_mode"
            )

    scope_mapping = _read_optional_sheet(path, "Scope Mapping")
    category_mapping = _read_optional_sheet(path, "Category Mapping")
    part_overrides = _read_optional_sheet(path, "Part Overrides")
    planned_basket = _read_optional_sheet(path, "Planned Basket")

    return ResolvedConfig(
        controls=controls,
        sources=sources,
        scope_mapping=scope_mapping,
        category_mapping=category_mapping,
        part_overrides=part_overrides,
        planned_basket=planned_basket,
        config_path=path,
    )


def resolve_dates(controls: ControlDefaults, latest_po_date: date) -> tuple[date, date]:
    base = controls.base_date or latest_po_date
    target = controls.target_date or (base + timedelta(days=365))
    if target < base:
        raise ValueError(
            f"target_date {target} is earlier than base_date {base}. "
            "Use a historical-index query for retrospective multipliers."
        )
    return base, target
