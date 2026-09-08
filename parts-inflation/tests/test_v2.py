"""Focused regression tests for the v2 production methodology."""

from __future__ import annotations

import json

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from parts_inflation.benchmarks import period_part_prices, tornqvist_index
from parts_inflation.classify import apply_cutoff_category_labels, apply_scope_and_category
from parts_inflation.clean import clean_po_lines
from parts_inflation.committed import build_committed_cost_indicator
from parts_inflation.config import load_config, write_config_workbook
from parts_inflation.repeat_sales import RepeatSalesResult
from parts_inflation.v2_model import _actual_matched_basket, build_bucket_weights, build_forecasts
from parts_inflation.v2_report import write_v2_outputs


@pytest.fixture
def cfg(tmp_path):
    path = tmp_path / "model_config.xlsx"
    write_config_workbook(path)
    return load_config(path)


def _raw_rows() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "P.O. Date": ["2024-01-01", "2025-01-01"],
            "Part Number": ["A1", "A1"],
            "Description 1": ["WIDGET", "WIDGET"],
            "Description 2": ["", ""],
            "Cost": [10.0, 11.0],
            "Qty Ordered": [5.0, 5.0],
            "Qty Received": [5.0, 0.0],
            "Extension (Qty Received)": [50.0, 0.0],
            "PO Value": [50.0, 55.0],
            "Bucket": ["1210AC", "1210AC"],
            "Description": ["Inventory", "Inventory"],
            "source_file": ["t.xlsx", "t.xlsx"],
            "source_sheet": ["Sheet2", "Sheet2"],
            "source_row_number": [2, 3],
            "raw_part_number": ["A1", "A1"],
            "raw_part_number_type": ["str", "str"],
        }
    )


def test_nonadjacent_benchmark_period_is_not_assigned_to_endpoint():
    daily = pd.DataFrame(
        {
            "PartKey": ["A", "A"],
            "po_date": pd.to_datetime(["2024-01-15", "2024-03-15"]),
            "price": [10.0, 20.0],
            "qty": [1.0, 1.0],
            "spend": [10.0, 20.0],
            "approved_category": ["Inventory", "Inventory"],
        }
    )
    assert tornqvist_index(period_part_prices(daily, "M")).empty


def test_backtest_holdout_retains_but_robustly_bounds_price_basis_extreme(cfg):
    daily = pd.DataFrame(
        {
            "PartKey": ["A", "B", "A", "B"],
            "po_date": pd.to_datetime(
                ["2024-01-01", "2024-01-01", "2025-01-01", "2025-01-01"]
            ),
            "price": [10.0, 10.0, 10_000.0, 11.0],
            "qty": [1.0] * 4,
            "spend": [10.0] * 4,
            "approved_category": ["Inventory"] * 4,
        }
    )
    actual = _actual_matched_basket(daily, pd.Timestamp("2024-01-01"), 12, cfg)
    assert actual["raw_actual_multiplier"].max() == pytest.approx(1000.0)
    assert actual["actual_multiplier"].max() <= cfg.controls.extreme_ratio_high
    assert actual.loc[actual["actual_extreme_flag"], "weight"].sum() < 0.5


def test_realized_and_committed_prices_are_separate(cfg):
    cleaned = clean_po_lines(_raw_rows(), cfg)
    classified, _ = apply_scope_and_category(cleaned, cfg)
    detail, summary = build_committed_cost_indicator(
        classified, pd.Timestamp("2025-01-01"), minimum_gap_days=30
    )
    assert cleaned.loc[0, "historical_unit_price"] == pytest.approx(10.0)
    assert pd.isna(cleaned.loc[1, "historical_unit_price"])
    assert cleaned.loc[1, "committed_unit_price"] == pytest.approx(11.0)
    assert detail.loc[0, "prior_realized_price"] == pytest.approx(10.0)
    overall = summary.loc[summary["category"].eq("Overall Direct Costs")].iloc[0]
    assert overall["matched_open_value_coverage"] == pytest.approx(1.0)
    assert overall["annualized_rate_signal"] == pytest.approx(0.10, abs=0.002)


def test_uom_override_changes_price_basis_without_changing_value(cfg):
    cfg.part_overrides = pd.DataFrame(
        [{"PartKey": "A1", "UoMAdjustmentFactor": 1000.0}]
    )
    cleaned = clean_po_lines(_raw_rows().iloc[[0]].copy(), cfg)
    classified, _ = apply_scope_and_category(cleaned, cfg)
    assert classified.loc[0, "historical_unit_price"] == pytest.approx(0.01)
    assert classified.loc[0, "qty_received"] == pytest.approx(5000.0)
    assert classified.loc[0, "realized_spend"] == pytest.approx(50.0)


def test_category_mode_uses_only_information_known_at_cutoff():
    frame = pd.DataFrame(
        {
            "comparison_entity_id": ["A", "A", "A"],
            "PartKey": ["A", "A", "A"],
            "in_direct_costs": [True, True, True],
            "effective_historical_date": pd.to_datetime(
                ["2024-01-01", "2025-01-01", "2025-02-01"]
            ),
            "direct_cost_category": [
                "Inventory", "Production Supplies", "Production Supplies"
            ],
        }
    )
    out = apply_cutoff_category_labels(frame, pd.Timestamp("2024-06-30"))
    assert set(out["approved_category"]) == {"Inventory"}


def test_fixed_basket_composite_uses_weighted_price_multipliers(cfg):
    months = list(pd.period_range("2021-01", periods=48, freq="M"))
    inv = np.full(48, np.log(1.10) / 12)
    prod = np.full(48, np.log(1.20) / 12)
    overall = 0.25 * inv + 0.75 * prod
    model = RepeatSalesResult(
        months=months,
        categories=["Inventory", "Production Supplies"],
        delta0=overall,
        u=np.vstack([inv - overall, prod - overall]),
        gamma0=0.0,
        kappa=np.zeros(2),
        pair_weights=np.ones(200),
        sigma=0.1,
        converged=True,
        n_pairs=200,
        n_pairs_used=200,
    )
    pairs = pd.DataFrame(
        {"approved_category": ["Inventory"] * 100 + ["Production Supplies"] * 100}
    )
    weights = pd.DataFrame(
        {
            "direct_cost_category": ["Inventory", "Production Supplies"],
            "weight": [0.25, 0.75],
        }
    )
    result = build_forecasts(
        model, pairs, weights, pd.Timestamp("2024-12-31"), cfg,
        forced_method="trailing_12m_mean",
    )
    year1 = result.composite_forecast.iloc[0]
    assert year1["annual_multiplier"] == pytest.approx(0.25 * 1.10 + 0.75 * 1.20)


def test_forecast_guardrail_caps_explosive_bucket_path(cfg):
    months = list(pd.period_range("2023-01", periods=24, freq="M"))
    explosive = np.full(24, 0.20)
    model = RepeatSalesResult(
        months=months, categories=["Inventory"], delta0=explosive,
        u=np.zeros((1, 24)), gamma0=0.0, kappa=np.zeros(1),
        pair_weights=np.ones(100), sigma=0.1, converged=True,
        n_pairs=100, n_pairs_used=100,
    )
    pairs = pd.DataFrame({"approved_category": ["Inventory"] * 100})
    weights = pd.DataFrame(
        {"direct_cost_category": ["Inventory"], "weight": [1.0]}
    )
    result = build_forecasts(
        model, pairs, weights, pd.Timestamp("2024-12-31"), cfg,
        forced_method="trailing_12m_mean",
    )
    bucket = result.bucket_forecast.iloc[0]
    assert bucket["annual_multiplier"] == pytest.approx(1.50)
    assert bool(bucket["forecast_cap_applied"])


def test_planned_basket_quantity_fallback(cfg):
    cfg.controls.composite_weighting = "planned_basket"
    cfg.planned_basket = pd.DataFrame(
        {"PartKey": ["A", "B"], "ExpectedQuantity": [10.0, 10.0]}
    )
    classified = pd.DataFrame(
        {
            "PartKey": ["A", "B"],
            "in_direct_costs": [True, True],
            "effective_historical_date": pd.to_datetime(["2024-01-01", "2024-01-01"]),
            "direct_cost_category": ["Inventory", "Production Supplies"],
            "historical_unit_price": [2.0, 6.0],
        }
    )
    weights = build_bucket_weights(classified, pd.Timestamp("2024-12-31"), cfg)
    got = weights.set_index("direct_cost_category")["weight"]
    assert got["Inventory"] == pytest.approx(0.25)
    assert got["Production Supplies"] == pytest.approx(0.75)


def test_output_package_is_complete_and_reopenable(tmp_path):
    tables = {
        "historical_monthly": pd.DataFrame(
            {
                "month": ["2024-01", "2024-02"],
                "month_end": pd.to_datetime(["2024-01-31", "2024-02-29"]),
                "series": ["Combined Direct Costs"] * 2,
                "index": [100.0, 101.0],
            }
        )
    }
    payload = {
        "summary": {"executive_metrics": {"Forecast year 1 inflation": 0.05}},
        "resolved_config": {"scope_mode": "direct_costs"},
        "tables": tables,
        "manifest": {"run_id": "test"},
    }
    paths = write_v2_outputs(tmp_path, payload)
    wb = load_workbook(paths["excel"], read_only=True)
    assert {"Executive Summary", "Historical Monthly", "Forecasts", "Methodology"}.issubset(
        wb.sheetnames
    )
    wb.close()
    manifest = json.loads(paths["manifest"].read_text(encoding="utf-8"))
    assert "inflation_report.xlsx" in manifest["output_hashes"]
    assert paths["summary"].exists()
